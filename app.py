#!/usr/bin/env python3
"""
Cortex Drive - an agentic cloud workspace prototype.

Users talk to an agent layer; the agent layer orchestrates the file system.

Agents
  Librarian  - proactive organisation (auto-tag, cluster, name folders)
  Researcher - semantic retrieval with citations
  Writer     - cross-document synthesis
  Analyst    - trigger-based workflows (invoice -> expense ledger, meeting -> action tracker)
  Watcher    - proactive insights (figures that changed, reversed decisions, deadlines, duplicates)

Works fully offline (TF-IDF retrieval + extractive generation).
Set ANTHROPIC_API_KEY to upgrade answers / synthesis to Claude, and
USE_EMBEDDINGS=1 (needs sentence-transformers) to add dense retrieval.
"""
import csv
import os
import re
import shutil
import threading
import warnings
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import numpy as np
from flask import Flask, jsonify, request, send_from_directory
from sklearn.cluster import KMeans
from sklearn.feature_extraction.text import ENGLISH_STOP_WORDS, TfidfVectorizer
from sklearn.metrics import silhouette_score
from sklearn.metrics.pairwise import cosine_similarity

warnings.filterwarnings("ignore")

BASE = Path(__file__).parent
WS = BASE / "workspace"
FILES_DIR = WS / "files"
OUT_DIR = WS / "_agent_outputs"
SEED = BASE / "sample_files" / "seed"
DEMO = BASE / "sample_files" / "demo_uploads"
EXPENSES_CSV = OUT_DIR / "expenses.csv"
ACTIONS_CSV = OUT_DIR / "action_items.csv"
EXP_HEADER = ["date", "vendor", "invoice_no", "category", "amount_usd", "source_file"]
ACT_HEADER = ["date", "meeting", "owner", "task", "source_file"]

app = Flask(__name__, static_folder="static", static_url_path="/static")
LOCK = threading.RLock()

DOCS = {}          # id -> doc dict
LOG = []           # activity feed
NAMES = set()      # person names (kept out of folder names)
VEC = CV = CM = None
IDS = []
CHUNKS = []
EMB = None
EMB_M = None

NOISE = {"meeting", "transcript", "attendees", "action", "date", "decision", "call",
         "review", "notes", "team", "today", "need", "needs", "going"}
STOP = set(ENGLISH_STOP_WORDS) | {"meeting", "meetings", "discussed", "discuss", "said",
                                  "tell", "show", "find", "document", "documents", "files", "file"}
SYN = {
    "budget": ["spend", "cost", "allocation", "funding", "approved"],
    "spend": ["budget", "cost", "spent"],
    "deadline": ["due", "date"],
    "requirement": ["must", "need", "require"],
    "lead": ["leads"],
    "notification": ["push", "sms"],
    "risk": ["concern", "issue"],
}
MONEY = re.compile(r"\$\s?(\d[\d,]*(?:\.\d+)?)\s?([kKmM])?")
MONTHS = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun",
                                          "jul", "aug", "sep", "oct", "nov", "dec"])}

if os.environ.get("USE_EMBEDDINGS") == "1":
    try:
        from sentence_transformers import SentenceTransformer
        EMB = SentenceTransformer("all-MiniLM-L6-v2")
    except Exception:
        EMB = None


# ----------------------------------------------------------------- utilities
def log(agent, msg, kind="info"):
    LOG.append({"t": datetime.now().strftime("%H:%M:%S"), "agent": agent, "msg": msg, "kind": kind})
    del LOG[:-300]


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def norm(w):
    return w[:-1] if len(w) > 3 and w.endswith("s") else w


def fmt_date(iso):
    d = date.fromisoformat(iso)
    return f"{d.strftime('%b')} {d.day}"


def money_vals(s):
    out = []
    for m in MONEY.finditer(s):
        v = float(m.group(1).replace(",", ""))
        suf = (m.group(2) or "").lower()
        out.append(v * (1000 if suf == "k" else 1_000_000 if suf == "m" else 1))
    return out


def usd(v):
    return f"${v:,.0f}" if float(v).is_integer() else f"${v:,.2f}"


def usd2(v):
    return f"${float(v):,.2f}"


def llm(system, user, max_tokens=1500):
    key = os.environ.get("ANTHROPIC_API_KEY")
    if not key:
        return None
    try:
        import anthropic
        client = anthropic.Anthropic(api_key=key)
        r = client.messages.create(model=os.environ.get("AGENT_MODEL", "claude-sonnet-5-5"),
                                   max_tokens=max_tokens, system=system,
                                   messages=[{"role": "user", "content": user}])
        return r.content[0].text
    except Exception as e:  # fall back to offline mode
        log("System", f"LLM unavailable ({e.__class__.__name__}) - using offline mode", "error")
        return None


# ------------------------------------------------------------------ ingestion
def extract_text(p: Path):
    ext = p.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)
    if ext == ".docx":
        import docx
        return "\n".join(par.text for par in docx.Document(str(p)).paragraphs)
    return p.read_text(errors="ignore")


def classify_type(name, text):
    t, n = text.lower(), name.lower()
    if "invoice" in n or ("invoice" in t[:300] and "total" in t):
        return "invoice"
    if n.endswith(".csv"):
        return "spreadsheet"
    if re.search(r"\b(attendees|minutes|transcript)\b", t[:500]) or "meeting" in n or "call" in n:
        return "meeting"
    if "policy" in n or "policy" in t[:200]:
        return "policy"
    return "notes"


def doc_date(text, path):
    m = re.search(r"^\s*(?:invoice )?date\s*:\s*(\d{4}-\d{2}-\d{2})", text, re.I | re.M)
    if m:
        try:
            return date.fromisoformat(m.group(1)).isoformat()
        except ValueError:
            pass
    return datetime.fromtimestamp(path.stat().st_mtime).date().isoformat()


INV_CATS = [
    ("Infrastructure", r"\b(cloud|hosting|servers?|aws|azure|compute|saas|software|licen[cs]e)"),
    ("Design & Creative", r"\b(design|studio|creative|brand|illustrat)"),
    ("Travel", r"\b(flights?|airline|hotel|travel|taxi|uber)"),
    ("Marketing", r"\b(advert|ads|campaign|influenc|sponsor)"),
    ("Office & Supplies", r"\b(office|supplies|furniture|stationery)"),
]


def parse_invoice(d):
    t = d["text"]

    def g(p):
        m = re.search(p, t, re.I | re.M)
        return m.group(1).strip() if m else ""

    total = g(r"\btotal(?: due| amount)?\s*:?\s*\$?\s*([\d,]+\.\d{2})")
    cat = "Other"
    for name, pat in INV_CATS:
        if re.search(pat, t.lower()):
            cat = name
            break
    return {
        "vendor": g(r"^(?:vendor|from|billed by)\s*:\s*(.+)$"),
        "invoice_no": g(r"invoice\s*(?:#|no\.?|number)\s*:?\s*([A-Z0-9-]+)"),
        "date": g(r"^(?:invoice )?date\s*:\s*(\d{4}-\d{2}-\d{2})") or d["date"],
        "amount": float(total.replace(",", "")) if total else 0.0,
        "category": cat,
    }


CUE_ORDER = [
    ("Action items", r"^\s*(?:action|todo)\b"),
    ("Decisions", r"\b(decision|decided|agreed|approved|going with|settled)\b"),
    ("Requirements & asks", r"\b(must|shall|need|needs|required?|requirements?|should|has to|have to|want|wants)\b"),
    ("Constraints, budget & deadlines", r"\b(deadline|due|budget|capped|cap|limit|gdpr|compliance|constraint)\b"),
    ("Risks & open questions", r"\b(risk|concern|worried|blocker|unclear|open question|not sure)\b|\?"),
]
REVERSAL = re.compile(r"\b(instead|no longer|drop|dropped|scrap|scrapped|pause|replace|replaced|postpone)\b", re.I)


def summarize(d):
    t = d["text"]
    lines = [l.strip() for l in t.splitlines() if l.strip()]
    if d["type"] == "invoice":
        inv = parse_invoice(d)
        return f"Invoice {inv['invoice_no']} from {inv['vendor']} for {usd2(inv['amount'])} ({inv['category']})"
    if d["type"] == "spreadsheet":
        return f"Spreadsheet, {max(len(lines) - 1, 0)} rows - columns: {lines[0][:80]}"
    title = re.sub(r"^#+\s*", "", lines[0])[:90] if lines else d["name"]
    body = [l for l in lines[1:] if len(l) > 40 and not re.match(r"^(date|attendees|effective)\b", l, re.I)]
    pick = next((re.sub(r"^[A-Z][a-z]+:\s*", "", l) for l in body if re.search(CUE_ORDER[1][1], l, re.I)), None)
    pick = pick or (re.sub(r"^[A-Z][a-z]+:\s*", "", body[0]) if body else "")
    return f"{title} - {pick}"[:230]


def ingest(path: Path, run_wf=True):
    text = extract_text(path)
    if not text.strip():
        log("Librarian", f"Skipped {path.name} (no readable text)", "error")
        return None
    did = slug(path.name)
    d = {"id": did, "name": path.name, "path": str(path), "text": text,
         "type": classify_type(path.name, text), "date": doc_date(text, path),
         "folder": DOCS.get(did, {}).get("folder"), "tags": []}
    d["summary"] = summarize(d)
    DOCS[did] = d
    log("Librarian", f"Read {path.name} - detected a {d['type']}", "info")
    if run_wf:
        run_workflows(d)
    return d


# ------------------------------------------------------------------ workflows
def replace_rows(path, header, src_idx, src, new_rows):
    rows = [header]
    if path.exists():
        with path.open(newline="") as f:
            rows = list(csv.reader(f)) or [header]
    rows = [r for i, r in enumerate(rows) if i == 0 or r[src_idx] != src] + new_rows
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as f:
        csv.writer(f).writerows(rows)


def read_csv(path):
    if not path.exists():
        return []
    with path.open(newline="") as f:
        return list(csv.DictReader(f))


def wf_invoice(d):
    inv = parse_invoice(d)
    replace_rows(EXPENSES_CSV, EXP_HEADER, 5, d["name"],
                 [[inv["date"], inv["vendor"], inv["invoice_no"], inv["category"],
                   f"{inv['amount']:.2f}", d["name"]]])
    return (f"Invoice detected -> extracted {inv['vendor']}, #{inv['invoice_no']}, {usd2(inv['amount'])}; "
            f"categorised as {inv['category']}; appended to expenses.csv")


def wf_actions(d):
    rows = []
    title = d["text"].splitlines()[0].strip()
    for line in d["text"].splitlines():
        m = re.match(r"^\s*action(?: items?)?\s*:\s*(.+)$", line, re.I)
        if m:
            task = m.group(1).strip()
            owner = re.match(r"^([A-Z][a-z]+)\b", task)
            rows.append([d["date"], title, owner.group(1) if owner else "", task, d["name"]])
    if not rows:
        return None
    replace_rows(ACTIONS_CSV, ACT_HEADER, 4, d["name"], rows)
    return f"Meeting transcript detected -> extracted {len(rows)} action items into action_items.csv"


WORKFLOWS = [
    {"id": "invoice-ledger", "agent": "Analyst", "name": "Invoice -> Expense Ledger",
     "trigger": "An invoice lands in the drive",
     "action": "Extract vendor, number, date and total, categorise it, append to expenses.csv",
     "match": lambda d: d["type"] == "invoice", "fn": wf_invoice},
    {"id": "meeting-actions", "agent": "Analyst", "name": "Meeting -> Action Tracker",
     "trigger": "A meeting transcript lands in the drive",
     "action": "Pull out every 'Action:' line with owner and date into action_items.csv",
     "match": lambda d: d["type"] == "meeting", "fn": wf_actions},
]


def run_workflows(d):
    for wf in WORKFLOWS:
        if wf["match"](d):
            try:
                msg = wf["fn"](d)
                if msg:
                    log(wf["agent"], msg, "workflow")
            except Exception as e:
                log("System", f"Workflow '{wf['name']}' failed: {e}", "error")


# ------------------------------------------------------------------- indexing
def make_chunks():
    out = []
    for d in DOCS.values():
        lines = [l.strip() for l in d["text"].splitlines() if l.strip()]
        buf = ""
        for l in lines:
            buf = (buf + "\n" + l).strip()
            if len(buf) >= 260:
                out.append({"doc": d["id"], "text": buf})
                buf = ""
        if buf:
            out.append({"doc": d["id"], "text": buf})
    return out


def pick_k(X):
    n = X.shape[0]
    best, best_k = -2, 2
    for k in range(2, min(6, n)):
        labels = KMeans(n_clusters=k, n_init=10, random_state=0).fit_predict(X)
        if len(set(labels)) < 2:
            continue
        s = silhouette_score(X, labels, metric="cosine")
        if s > best:
            best, best_k = s, k
    return best_k


def reindex(full=False):
    """full=True re-clusters everything; otherwise only new files are filed (existing folders stay put)."""
    global VEC, CV, CM, IDS, CHUNKS, EMB_M
    IDS = list(DOCS)
    CHUNKS = []
    if not IDS:
        return
    NAMES.clear()
    for d in DOCS.values():
        NAMES.update(m.group(1).lower() for m in re.finditer(r"^([A-Z][a-z]+):", d["text"], re.M))
        m = re.search(r"^Attendees:\s*(.+)$", d["text"], re.M)
        if m:
            NAMES.update(w.lower() for w in re.findall(r"([A-Z][a-z]+)(?:\s*\(|,|$)", m.group(1)))

    VEC = TfidfVectorizer(stop_words=list(ENGLISH_STOP_WORDS | NOISE), ngram_range=(1, 2), sublinear_tf=True)
    X = VEC.fit_transform([DOCS[i]["text"] for i in IDS])
    terms = VEC.get_feature_names_out()

    def good_terms(vec_row, limit):
        out = []
        for j in np.asarray(vec_row).ravel().argsort()[::-1]:
            t = terms[j]
            if " " in t or not t.isalpha() or len(t) < 4 or t in NAMES or t in out:
                continue
            out.append(t)
            if len(out) == limit:
                break
        return out

    # ---- filing: invoices get their own folder, everything else is clustered by content
    prev_f = {i: DOCS[i].get("folder") for i in IDS}

    def cluster_name(idxs):
        docs_ = [DOCS[IDS[i]] for i in idxs]
        if all(d["type"] == "invoice" for d in docs_):
            return "Finance - Invoices"
        titles = [re.sub(r"^#+\s*", "", d["text"].splitlines()[0].strip()) for d in docs_ if d["type"] != "spreadsheet"]
        heads = [re.split(r"\s[-–—]\s", t)[0].split() for t in titles]
        if len(heads) == 1:
            return titles[0] if len(titles[0]) <= 45 else " ".join(heads[0])
        pref = []
        for ws in zip(*heads):
            if len(set(ws)) != 1:
                break
            pref.append(ws[0])
        if pref:
            return " ".join(pref)
        tt = good_terms(np.asarray(X[idxs].mean(axis=0)), 2)
        return " & ".join(w.title() for w in tt) or "Miscellaneous"

    def full_cluster():
        out = {}
        inv = [k for k, i in enumerate(IDS) if DOCS[i]["type"] == "invoice"]
        sheets = [k for k, i in enumerate(IDS) if DOCS[i]["type"] == "spreadsheet"]
        rest = [k for k, i in enumerate(IDS) if DOCS[i]["type"] not in ("invoice", "spreadsheet")]
        for k in inv:
            out[IDS[k]] = "Finance - Invoices"
        if len(rest) >= 4:
            Xr = X[rest]
            labels = KMeans(n_clusters=pick_k(Xr), n_init=10, random_state=0).fit_predict(Xr)
        else:
            labels = np.zeros(len(rest), dtype=int)
        clusters = defaultdict(list)
        for pos, lab in enumerate(labels):
            clusters[int(lab)].append(rest[pos])
        old = defaultdict(set)
        for did, f in prev_f.items():
            if f:
                old[f].add(did)
        used = {"Finance - Invoices"} if inv else set()
        for lab, idxs in sorted(clusters.items(), key=lambda kv: -len(kv[1])):
            members = {IDS[i] for i in idxs}
            best, best_ov = None, 0.0
            for name, om in old.items():
                ov = len(members & om) / len(members | om)
                if ov > best_ov and name not in used:
                    best, best_ov = name, ov
            fname = best if best and best_ov >= 0.5 else cluster_name(idxs)
            base, k2 = fname, 2
            while fname in used:
                fname, k2 = f"{base} ({k2})", k2 + 1
            used.add(fname)
            for i in idxs:
                out[IDS[i]] = fname
        # spreadsheets join the folder whose documents they resemble most
        for k in sheets:
            best, bs = None, 0.0
            for lab, idxs in clusters.items():
                s = cosine_similarity(X[k], np.asarray(X[idxs].mean(axis=0)))[0, 0]
                if s > bs:
                    best, bs = out[IDS[idxs[0]]], s
            if best:
                out[IDS[k]] = best
            else:
                nm = cluster_name([k]) if False else DOCS[IDS[k]]["name"]
                out[IDS[k]] = nm
        return out

    if full or not any(prev_f.values()):
        assign = full_cluster()
    else:
        assign = {i: f for i, f in prev_f.items() if f}
        groups = defaultdict(list)
        for k, i in enumerate(IDS):
            if i in assign and assign[i] != "Finance - Invoices":
                groups[assign[i]].append(k)
        cents = {f: np.asarray(X[ks].mean(axis=0)) for f, ks in groups.items()}
        for k, i in enumerate(IDS):
            if i in assign:
                continue
            if DOCS[i]["type"] == "invoice":
                assign[i] = "Finance - Invoices"
                continue
            best, bs = None, 0.0
            for f, c in cents.items():
                s = cosine_similarity(X[k], c)[0, 0]
                if s > bs:
                    best, bs = f, s
            if best and bs >= 0.08:
                assign[i] = best
            else:
                nm = base = cluster_name([k])
                n2 = 2
                while nm in set(assign.values()):
                    nm, n2 = f"{base} ({n2})", n2 + 1
                assign[i] = nm

    labels_t = {"invoice": "Invoice", "meeting": "Meeting", "spreadsheet": "Spreadsheet",
                "policy": "Policy", "notes": "Notes"}
    for i, did in enumerate(IDS):
        d = DOCS[did]
        prev = d.get("folder")
        d["folder"] = assign[did]
        if d["type"] == "invoice":
            inv = parse_invoice(d)
            d["tags"] = ["Invoice", inv["category"], inv["vendor"]]
        else:
            d["tags"] = [labels_t[d["type"]]] + good_terms(X[i].toarray(), 3)
        if prev is None:
            log("Librarian", f"Filed {d['name']} -> '{d['folder']}' (tags: {', '.join(d['tags'][1:])})", "organize")
        elif prev != d["folder"]:
            log("Librarian", f"Moved {d['name']}: '{prev}' -> '{d['folder']}'", "organize")

    # chunk index for retrieval
    CHUNKS = make_chunks()
    CV = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), sublinear_tf=True)
    CM = CV.fit_transform([c["text"] for c in CHUNKS])
    EMB_M = EMB.encode([c["text"] for c in CHUNKS], normalize_embeddings=True) if EMB else None


def expand(q):
    toks = [norm(w) for w in re.findall(r"[a-z0-9$]+", q.lower())]
    extra = [s for t in toks for s in SYN.get(t, [])]
    return q + " " + " ".join(extra)


def retrieve(q, k=6, per_doc=2, restrict=None):
    if not CHUNKS:
        return []
    sims = cosine_similarity(CV.transform([expand(q)]), CM)[0]
    if EMB is not None:
        sims = 0.5 * sims + 0.5 * (EMB_M @ EMB.encode([q], normalize_embeddings=True)[0])
    out, cnt = [], defaultdict(int)
    for i in sims.argsort()[::-1]:
        if sims[i] <= 0:
            break
        c = CHUNKS[i]
        if (restrict and c["doc"] not in restrict) or cnt[c["doc"]] >= per_doc:
            continue
        out.append({**c, "score": float(sims[i])})
        cnt[c["doc"]] += 1
        if len(out) >= k:
            break
    return out


# -------------------------------------------------------------------- insights
def budget_timeline(restrict=None):
    """Find figures attached to the word 'budget' that changed across documents in the same folder."""
    groups = defaultdict(list)  # (folder, quarter) -> [(date, amount, doc)]
    for d in DOCS.values():
        if d["type"] not in ("meeting", "notes") or (restrict and d["id"] not in restrict):
            continue
        for line in d["text"].splitlines():
            bm = re.search(r"budget", line, re.I)
            if not bm:
                continue
            ms = list(MONEY.finditer(line))
            if not ms:
                continue
            m = min(ms, key=lambda m: abs(m.start() - bm.start()))
            amt = money_vals(m.group(0))[0]
            q = re.search(r"\bQ[1-4]\b", line)
            groups[(d["folder"], q.group(0) if q else "this project")].append((d["date"], amt, d))
    out = []
    for (folder, q), items in groups.items():
        items.sort(key=lambda x: x[0])
        steps = []
        for dt, amt, d in items:
            if not steps or steps[-1][1] != amt:
                steps.append((dt, amt, d))
        if len(steps) >= 2:
            out.append((q, steps))
    return out


def reversals(restrict=None):
    docs = sorted([d for d in DOCS.values() if d["type"] in ("meeting", "notes")
                   and (not restrict or d["id"] in restrict)], key=lambda d: d["date"])
    found = []
    for d in docs:
        for line in d["text"].splitlines():
            line = line.strip()
            if not REVERSAL.search(line) or len(line) < 30:
                continue
            toks = {norm(w) for w in re.findall(r"[a-z0-9]+", line.lower())
                    if w not in STOP and len(w) >= 3 and w not in NAMES and not REVERSAL.match(w)}
            best, best_ov = None, 1
            for e in docs:
                if e["date"] >= d["date"] or e["folder"] != d["folder"]:
                    continue
                for l2 in e["text"].splitlines():
                    t2 = {norm(w) for w in re.findall(r"[a-z0-9]+", l2.lower())}
                    ov = len(toks & t2 - {"will", "use"})
                    if ov > best_ov and not REVERSAL.search(l2):
                        best, best_ov = (e, l2.strip()), ov
            if best:
                found.append((d, line, best[0], best[1]))
    return found


def deadlines(restrict=None):
    today = date.today()
    items, seen = [], set()
    pat = re.compile(r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+(\d{1,2})\b")
    for d in DOCS.values():
        if d["type"] in ("invoice", "spreadsheet") or (restrict and d["id"] not in restrict):
            continue
        for line in d["text"].splitlines():
            m = pat.search(line)
            if not m or not re.search(r"\b(by|before|due|deadline|until|ready|submit)\b", line, re.I):
                continue
            try:
                dt = date(int(d["date"][:4]), MONTHS[m.group(1).lower()], int(m.group(2)))
            except ValueError:
                continue
            txt = re.sub(r"^[A-Z][a-z]+:\s*", "", line.strip())
            key = (dt, txt[:40])
            if key in seen:
                continue
            seen.add(key)
            items.append((dt, txt, d))
    return sorted([i for i in items if i[0] >= today], key=lambda x: x[0]), today


def duplicates():
    if VEC is None or len(IDS) < 2:
        return []
    X = VEC.transform([DOCS[i]["text"] for i in IDS])
    S = cosine_similarity(X)
    out = []
    for i in range(len(IDS)):
        for j in range(i + 1, len(IDS)):
            if S[i, j] > 0.85:
                out.append((IDS[i], IDS[j], float(S[i, j])))
    return out


def insights(restrict=None):
    cards = []
    for q, steps in budget_timeline(restrict):
        trail = " → ".join(f"{usd(a)} ({fmt_date(dt)})" for dt, a, _ in steps)
        cards.append({"kind": "figure", "title": ("Project" if q == "this project" else q) + " budget figure changed",
                      "detail": trail, "sources": [d["id"] for _, _, d in steps]})
    for d, line, e, l2 in reversals(restrict):
        cards.append({"kind": "reversal", "title": "Change of direction detected",
                      "detail": f"{fmt_date(e['date'])} ({e['name']}): \"{re.sub(r'^[A-Z][a-z]+:\\s*', '', l2)}\"\n"
                                f"then {fmt_date(d['date'])} ({d['name']}): \"{re.sub(r'^[A-Z][a-z]+:\\s*', '', line)}\"",
                      "sources": [e["id"], d["id"]]})
    dl, today = deadlines(restrict)
    for dt, txt, d in dl[:5]:
        left = (dt - today).days
        cards.append({"kind": "deadline", "title": f"Due {fmt_date(dt.isoformat())} - in {left} day{'s' if left != 1 else ''}",
                      "detail": txt, "sources": [d["id"]]})
    for a, b, s in duplicates():
        cards.append({"kind": "duplicate", "title": "Near-duplicate files",
                      "detail": f"{DOCS[a]['name']} and {DOCS[b]['name']} are {int(s * 100)}% similar",
                      "sources": [a, b]})
    return cards


# --------------------------------------------------------------------- agents
def split_sentences(text):
    out = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        out.extend(p.strip() for p in re.split(r"(?<=[.!?])\s+(?=[A-Z])", line) if p.strip())
    return out


def ask_agent(q, trace):
    trace.append("Router -> Researcher agent (semantic retrieval)")
    hits = retrieve(q, k=7, per_doc=2)
    if not hits:
        return "I couldn't find anything relevant in your drive.", []
    ndocs = len({h["doc"] for h in hits})
    trace.append(f"Retrieved {len(hits)} passages from {ndocs} files")
    qtoks = {norm(w) for w in re.findall(r"[a-z0-9$]+", q.lower())} - STOP
    syn = {norm(s) for t in qtoks for s in SYN.get(t, [])} - qtoks
    money_q = bool(re.search(r"budget|cost|spend|spent|price|how much|total|amount|\$", q, re.I))

    # LLM path
    if os.environ.get("ANTHROPIC_API_KEY"):
        order = []
        for h in hits:
            if h["doc"] not in order:
                order.append(h["doc"])
        ctx = "\n\n".join(f"[{order.index(h['doc']) + 1}] {DOCS[h['doc']]['name']} ({DOCS[h['doc']]['date']}):\n{h['text']}"
                          for h in hits)
        ans = llm("You answer questions about the user's files. Use only the numbered sources, cite them like [1], "
                  "mention when figures changed over time, and be concise.", f"Question: {q}\n\nSources:\n{ctx}")
        if ans:
            trace.append("Claude composed a cited answer")
            return ans, [mk_source(i + 1, DOCS[did], next(h["text"] for h in hits if h["doc"] == did))
                         for i, did in enumerate(order)]

    cands = []
    for h in hits:
        for s in split_sentences(h["text"]):
            toks = {norm(w) for w in re.findall(r"[a-z0-9$]+", s.lower())}
            sc = len(toks & qtoks) + 0.5 * len(toks & syn)
            if sc > 0 and money_q and MONEY.search(s):
                sc += 1.5
            if sc > 0:
                cands.append((sc, h["doc"], s))
    if not cands:
        return "I found related files but nothing that directly answers that.", []
    best = max(c[0] for c in cands)
    keep, per = [], defaultdict(int)
    for sc, did, s in sorted(cands, key=lambda c: -c[0]):
        if sc < max(1.0, 0.45 * best) or per[did] >= 2 or len(keep) >= 6:
            continue
        keep.append((did, s))
        per[did] += 1
    keep.sort(key=lambda x: DOCS[x[0]]["date"])
    order = []
    for did, _ in keep:
        if did not in order:
            order.append(did)
    trace.append(f"Ranked sentences and kept the {len(keep)} most relevant across {len(order)} files")
    lines = [f"- **{fmt_date(DOCS[did]['date'])}** - {s} [{order.index(did) + 1}]" for did, s in keep]
    note = ""
    if money_q:
        tl = [t for t in budget_timeline(set(order)) if len(t[1]) >= 2]
        if tl:
            q_, steps = tl[0]
            note = ("\n\n**Heads up - this figure moved over time:** "
                    + " → ".join(f"{usd(a)} ({fmt_date(dt)})" for dt, a, _ in steps))
            trace.append("Watcher agent flagged a figure that changed between files")
    ans = f"Here's what I found across **{len(order)} files**:\n\n" + "\n".join(lines) + note
    srcs = [mk_source(i + 1, DOCS[did], next((s for d2, s in keep if d2 == did), "")) for i, did in enumerate(order)]
    return ans, srcs


def mk_source(n, d, snippet):
    return {"n": n, "id": d["id"], "name": d["name"], "date": d["date"], "folder": d["folder"],
            "snippet": snippet[:220]}


def title_for(docs, instruction):
    prefixes = {re.split(r"\s[-—–]\s", d["text"].splitlines()[0].strip())[0] for d in docs}
    if len(prefixes) == 1:
        return f"{prefixes.pop()} - Unified Requirements" if "requirement" in instruction.lower() else prefixes.pop() + " - Synthesis"
    return "Synthesis: " + instruction[:60]


def synth_offline(instruction, docs):
    src_no = {d["id"]: i + 1 for i, d in enumerate(docs)}
    buckets = {name: [] for name, _ in CUE_ORDER}
    for d in docs:
        if d["type"] in ("invoice", "spreadsheet"):
            continue
        for line in d["text"].splitlines():
            line = line.strip()
            if len(line) < 25 or re.match(r"^(date|attendees|meeting|effective)\b.*:", line, re.I):
                continue
            for name, pat in CUE_ORDER:
                if re.search(pat, line, re.I):
                    content = re.sub(r"^(?:action\s*:\s*)", "", line, flags=re.I)
                    content = re.sub(r"^[A-Z][a-z]+:\s*", "", content).strip()
                    content = re.sub(r"^(decision|open question)\s*:\s*", "", content, flags=re.I)
                    content = content[0].upper() + content[1:]
                    toks = {norm(w) for w in re.findall(r"[a-z0-9]+", content.lower()) if w not in STOP}
                    dup = next((b for b in buckets[name]
                                if len(toks & b["toks"]) / max(len(toks | b["toks"]), 1) >= 0.6), None)
                    if dup:
                        dup["srcs"].add(src_no[d["id"]])
                    else:
                        buckets[name].append({"text": content if content.endswith((".", "?")) else content + ".",
                                              "toks": toks, "srcs": {src_no[d["id"]]}})
                    break
    title = title_for(docs, instruction)
    md = [f"# {title}", f"*Drafted by the Writer agent from {len(docs)} files on {date.today().isoformat()}*", ""]
    md.append("## Overview")
    md.append(f"This document consolidates {len(docs)} source files ({docs[0]['name']}"
              + (f" ... {docs[-1]['name']}" if len(docs) > 1 else "") + "). Every line below links back to its source.")
    md.append("")
    for name, _ in CUE_ORDER:
        if buckets[name]:
            md.append(f"## {name}")
            for b in buckets[name]:
                md.append(f"- {b['text']} [{', '.join(str(n) for n in sorted(b['srcs']))}]")
            md.append("")
    ids = {d["id"] for d in docs}
    conf = []
    for d, line, e, l2 in reversals(ids):
        conf.append(f"- **Direction changed:** \"{re.sub(r'^[A-Z][a-z]+:\\s*', '', l2)}\" [{src_no[e['id']]}] was later revised: "
                    f"\"{re.sub(r'^[A-Z][a-z]+:\\s*', '', line)}\" [{src_no[d['id']]}]. Resolve before building.")
    for q, steps in budget_timeline(ids):
        conf.append(f"- **{'Project' if q == 'this project' else q} budget changed:** " + " → ".join(f"{usd(a)} [{src_no[d['id']]}]" for _, a, d in steps))
    if conf:
        md += ["## Conflicts & changes across sources"] + conf + [""]
    md.append("## Sources")
    md += [f"{i + 1}. {d['name']} ({fmt_date(d['date'])})" for i, d in enumerate(docs)]
    return "\n".join(md)


def synth_agent(msg, selected, trace):
    trace.append("Router -> Writer agent (cross-document synthesis)")
    if selected:
        ids = [i for i in selected if i in DOCS]
        trace.append(f"Using the {len(ids)} files you selected")
    else:
        hits = retrieve(msg, k=30, per_doc=3)
        score = defaultdict(float)
        for h in hits:
            score[h["doc"]] += h["score"]
        top = sorted(score.items(), key=lambda kv: -kv[1])
        ids = [d for d, s in top if s >= 0.4 * top[0][1] and DOCS[d]["type"] not in ("invoice", "spreadsheet")][:5] if top else []
        if ids:
            folder = DOCS[ids[0]]["folder"]
            extra = [d["id"] for d in DOCS.values() if d["folder"] == folder
                     and d["type"] not in ("invoice", "spreadsheet") and d["id"] not in ids]
            if extra:
                ids += extra
                trace.append(f"Pulled in {len(extra)} more file{'s' if len(extra) != 1 else ''} from the same project folder '{folder}'")
        trace.append(f"Auto-selected {len(ids)} relevant files: " + ", ".join(DOCS[i]["name"] for i in ids))
    if not ids:
        return "I couldn't work out which files to use - select some on the left and try again.", None, []
    docs = sorted([DOCS[i] for i in ids], key=lambda d: d["date"])
    trace.append("Reading every file and extracting requirements, decisions, constraints, risks and actions")
    md = None
    if os.environ.get("ANTHROPIC_API_KEY"):
        body = "\n\n".join(f"### [{i + 1}] {d['name']} ({d['date']})\n{d['text'][:6000]}" for i, d in enumerate(docs))
        md = llm("You are a careful technical writer. Produce a Markdown document from the numbered sources. "
                 "Cite sources inline like [1]. End with a 'Conflicts & changes across sources' section listing any "
                 "contradictions or decisions that were reversed, and a 'Sources' list.",
                 f"Instruction: {msg}\n\n{body}", max_tokens=3000)
        if md:
            trace.append("Claude drafted the document")
    if not md:
        md = synth_offline(msg, docs)
    trace.append("Cross-checked sources for conflicts and added citations")
    title = md.splitlines()[0].lstrip("# ").strip()
    srcs = [mk_source(i + 1, d, d["summary"]) for i, d in enumerate(docs)]
    ans = (f"I read **{len(docs)} files** ({', '.join(d['name'] for d in docs)}) and drafted "
           f"**{title}**. Review it below, then save it into your drive.")
    return ans, {"title": title, "markdown": md}, srcs


def expenses_agent(trace):
    trace.append("Router -> Analyst agent (expense ledger)")
    rows = read_csv(EXPENSES_CSV)
    if not rows:
        return "The expense ledger is empty - drop an invoice into the drive and the Analyst will log it automatically.", []
    by = defaultdict(float)
    for r in rows:
        by[r["category"]] += float(r["amount_usd"])
    total = sum(by.values())
    trace.append(f"Summed {len(rows)} invoices from expenses.csv")
    ans = f"**{len(rows)} invoices logged - total {usd2(total)}**\n\n" + "\n".join(
        f"- {c}: {usd2(v)} ({v / total:.0%})" for c, v in sorted(by.items(), key=lambda kv: -kv[1]))
    big = max(rows, key=lambda r: float(r["amount_usd"]))
    ans += f"\n\nLargest: {big['vendor']} #{big['invoice_no']} at {usd2(big['amount_usd'])}."
    return ans, [{"label": c, "value": v} for c, v in sorted(by.items(), key=lambda kv: -kv[1])]


def organize_agent(trace):
    trace.append("Router -> Librarian agent (proactive organisation)")
    before = {i: d["folder"] for i, d in DOCS.items()}
    reindex(True)
    moved = [i for i in DOCS if before.get(i) != DOCS[i]["folder"]]
    trace.append(f"Re-clustered {len(DOCS)} files by content; {len(moved)} moved")
    folders = defaultdict(list)
    for d in DOCS.values():
        folders[d["folder"]].append(d)
    lines = [f"- **{f}** - {len(ds)} file{'s' if len(ds) != 1 else ''}: " + ", ".join(d["name"] for d in ds[:4]) + (" ..." if len(ds) > 4 else "")
             for f, ds in folders.items()]
    return f"Your drive is organised into **{len(folders)} folders** (no manual filing):\n\n" + "\n".join(lines)


def insights_agent(trace):
    trace.append("Router -> Watcher agent (cross-document analysis)")
    cards = insights()
    trace.append(f"Scanned {len(DOCS)} files for changed figures, reversed decisions, deadlines and duplicates")
    if not cards:
        return "Nothing unusual found - no conflicting figures, reversed decisions or duplicates."
    lines = []
    for c in cards:
        lines.append(f"**{c['title']}**\n" + "\n".join(f"> {l}" for l in c["detail"].split("\n")))
    return "\n\n".join(lines)


def route(msg, selected):
    m = msg.lower()
    if re.search(r"\b(organi[sz]e|re-?organi[sz]e|clean up|restructure)\b", m) or re.search(r"sort (my|the) (drive|files)", m):
        return "organize"
    if re.search(r"\b(invoices?|expenses?|ledger|bills?)\b", m) and \
            re.search(r"\b(total|how much|summary|summari[sz]e|breakdown|categor\w*|spent|so far)\b", m):
        return "expenses"
    if re.search(r"\b(conflict|contradict|inconsisten\w*|changed|change their mind|changed their mind|duplicates?|insights?|"
                 r"anything (odd|off|new)|what.s new|reversal|flag)\b", m) or "mind" in m:
        return "insights"
    if re.search(r"\b(draft|write|generate|compile|create|produce|prepare|summari[sz]e)\b", m) and \
            (selected or re.search(r"\b(document|brief|report|plan|spec|requirements?|memo|proposal|summary)\b", m)):
        return "synth"
    return "ask"


def run_agent(msg, selected):
    trace = []
    intent = route(msg, selected)
    out = {"intent": intent, "sources": [], "doc": None, "chart": None}
    if intent == "ask":
        out["answer"], out["sources"] = ask_agent(msg, trace)
    elif intent == "synth":
        out["answer"], out["doc"], out["sources"] = synth_agent(msg, selected, trace)
        if out["doc"]:
            log("Writer", f"Drafted '{out['doc']['title']}' from {len(out['sources'])} files", "synth")
    elif intent == "expenses":
        out["answer"], out["chart"] = expenses_agent(trace)
    elif intent == "organize":
        out["answer"] = organize_agent(trace)
    else:
        out["answer"] = insights_agent(trace)
    out["trace"] = trace
    if intent == "ask" and out["sources"]:
        log("Researcher", f"Answered a question using {len(out['sources'])} files", "ask")
    return out


# -------------------------------------------------------------- state control
def load_dir(folder, run_wf):
    for p in sorted(folder.iterdir()):
        if p.is_file() and not p.name.startswith("."):
            ingest(p, run_wf)


def reset_demo():
    with LOCK:
        shutil.rmtree(WS, ignore_errors=True)
        FILES_DIR.mkdir(parents=True, exist_ok=True)
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        DOCS.clear()
        LOG.clear()
        for p in sorted(SEED.iterdir()):
            shutil.copy(p, FILES_DIR / p.name)
        log("System", "Demo workspace created", "info")
        load_dir(FILES_DIR, True)
        reindex()


def startup():
    if FILES_DIR.exists() and any(FILES_DIR.iterdir()):
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        load_dir(FILES_DIR, False)
        reindex()
    else:
        reset_demo()


# ----------------------------------------------------------------------- API
@app.get("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


@app.get("/api/status")
def status():
    return jsonify({"docs": len(DOCS), "llm": bool(os.environ.get("ANTHROPIC_API_KEY")),
                    "embeddings": EMB is not None,
                    "demo_left": len([p for p in DEMO.iterdir() if p.is_file() and not (FILES_DIR / p.name).exists()])})


@app.get("/api/files")
def files():
    folders = defaultdict(list)
    for d in DOCS.values():
        folders[d["folder"]].append({k: d[k] for k in ("id", "name", "type", "tags", "summary", "date")})
    return jsonify({"folders": [{"name": n, "files": sorted(fs, key=lambda f: f["date"])} for n, fs in
                                sorted(folders.items(), key=lambda kv: -len(kv[1]))],
                    "workflows": [{k: w[k] for k in ("id", "name", "trigger", "action", "agent")} for w in WORKFLOWS]})


@app.get("/api/file/<did>")
def file_preview(did):
    d = DOCS.get(did)
    if not d:
        return jsonify({"error": "not found"}), 404
    return jsonify({"name": d["name"], "text": d["text"][:8000], "folder": d["folder"], "tags": d["tags"]})


@app.get("/api/activity")
def activity():
    return jsonify(list(reversed(LOG[-80:])))


@app.get("/api/insights")
def api_insights():
    return jsonify(insights())


@app.get("/api/outputs")
def outputs():
    return jsonify({"expenses": read_csv(EXPENSES_CSV), "actions": read_csv(ACTIONS_CSV)})


@app.post("/api/upload")
def upload():
    with LOCK:
        start = len(LOG)
        names, errors = [], []
        uploaded = [f for f in request.files.getlist("files") if f and f.filename]
        if not uploaded:
            return jsonify({"error": "No files received. Select one or more files and try again."}), 400

        FILES_DIR.mkdir(parents=True, exist_ok=True)
        for f in uploaded:
            name = Path(f.filename).name
            try:
                p = FILES_DIR / name
                f.save(p)
                if ingest(p, True):
                    names.append(name)
                else:
                    errors.append({"name": name, "error": "No readable text was found in this file."})
            except Exception as exc:
                log("System", f"Could not process {name}: {exc.__class__.__name__}", "error")
                errors.append({"name": name, "error": f"Processing failed ({exc.__class__.__name__}). Check the file type and server terminal."})

        try:
            reindex()
        except Exception as exc:
            log("System", f"Search index update failed: {exc.__class__.__name__}", "error")
            errors.append({"name": "search index", "error": f"Index update failed ({exc.__class__.__name__}). Check the server terminal."})

        return jsonify({"ingested": names, "errors": errors, "events": LOG[start:]})


@app.post("/api/demo-drop")
def demo_drop():
    with LOCK:
        pending = [p for p in sorted(DEMO.iterdir()) if p.is_file() and not (FILES_DIR / p.name).exists()]
        if not pending:
            return jsonify({"done": True, "events": []})
        start = len(LOG)
        p = pending[0]
        shutil.copy(p, FILES_DIR / p.name)
        ingest(FILES_DIR / p.name, True)
        reindex()
        return jsonify({"done": False, "name": p.name, "events": LOG[start:]})


@app.post("/api/reset")
def reset():
    reset_demo()
    return jsonify({"ok": True})


@app.post("/api/agent")
def agent():
    body = request.get_json(force=True)
    with LOCK:
        return jsonify(run_agent(body.get("message", ""), body.get("selected", [])))


@app.post("/api/save")
def save():
    body = request.get_json(force=True)
    with LOCK:
        name = slug(body.get("title", "synthesis"))[:60] + ".md"
        p = FILES_DIR / name
        p.write_text(body.get("markdown", ""))
        ingest(p, True)
        reindex()
        log("Writer", f"Saved {name} into the drive", "synth")
        return jsonify({"saved": name})


startup()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5050)), debug=False)
